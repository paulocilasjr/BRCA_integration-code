from __future__ import annotations

from pathlib import Path

import matplotlib
import pytest
from openpyxl import Workbook

from brca_integration import pipeline
from brca_integration.tables.sup_table_18_19 import (
    _eve_points,
    cap_score,
    classify_final_points,
)
from brca_integration.validation import ValidationError, validate_generated_workbook


def _workbook(path: Path, sheets: tuple[str, ...]) -> None:
    workbook = Workbook()
    workbook.active.title = sheets[0]
    for sheet in sheets[1:]:
        workbook.create_sheet(sheet)
    workbook.save(path)


def test_eve_point_mapping() -> None:
    assert _eve_points("Benign") == -1
    assert _eve_points(" pathogenic ") == 1
    assert _eve_points("Uncertain") == 0
    assert _eve_points(None) == 0


@pytest.mark.parametrize(
    ("points", "expected"),
    [
        (-8, "B"),
        (-7, "B"),
        (-6, "LB"),
        (-2, "LB"),
        (-1, "VUS"),
        (5, "VUS"),
        (6, "LP"),
        (9, "LP"),
        (10, "P"),
    ],
)
def test_final_classification_thresholds(points: int, expected: str) -> None:
    assert classify_final_points(points) == expected


def test_functional_score_cap() -> None:
    assert cap_score(-20) == -4
    assert cap_score(-3) == -3
    assert cap_score(3) == 3
    assert cap_score(20) == 4


def test_figure_backend_is_headless() -> None:
    assert matplotlib.get_backend().lower() == "agg"


def test_generated_workbook_rejects_partial_file(tmp_path: Path) -> None:
    path = tmp_path / "partial.xlsx"
    _workbook(path, ("Sup Table 7", "Sup Table 8"))
    with pytest.raises(ValidationError, match="missing sheets"):
        validate_generated_workbook(path)


def test_failed_build_preserves_existing_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.xlsx"
    eve = tmp_path / "eve.xlsx"
    other = tmp_path / "other.xlsx"
    output = tmp_path / "published.xlsx"
    _workbook(source, tuple(f"Sup Table {number}" for number in range(1, 7)))
    _workbook(eve, ("README", "BRCA1", "BRCA2"))
    _workbook(other, ("BRCA1", "BRCA2"))
    output.write_bytes(b"previous-valid-output")

    def fail_build(**kwargs) -> None:
        Path(kwargs["output_workbook"]).write_bytes(b"partial")
        raise RuntimeError("simulated failure")

    monkeypatch.setattr(pipeline, "_build_supplementary_workbook", fail_build)
    with pytest.raises(RuntimeError, match="simulated failure"):
        pipeline.build_supplementary_workbook(
            input_workbook=source,
            output_workbook=output,
            eve_workbook=eve,
            other_points_workbook=other,
            figure_prefix=tmp_path / "figure" / "supp_fig2",
        )

    assert output.read_bytes() == b"previous-valid-output"
    assert not list(tmp_path.glob(".*.tmp.xlsx"))
